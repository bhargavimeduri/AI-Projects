"""Fix the 3 remaining [Image Placeholder] entries in Section 4.10."""

import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, RGBColor

INPUT = r"C:\Users\91630\OneDrive\Desktop\Bhargavi AIB\epaib-batch05-group4\Output Submission\Challenges_Issues_and_Resolutions_V2.docx"

doc = Document(INPUT)

remaining_map = {
    "Tiger walking alongside water body with its reflection visible":
        "Reference Image: 002655.jpg — Test Data 13\n(Tiger near water with reflection — combines water reflection suppression + multi-tiger risk. Pipeline applies water guard at grid and remainder scan levels.)",

    "Tiger partially occluded by tall grass, bamboo, or tree branches":
        "Reference Image: 002711.jpg — Test Data 13\n(Tiger partially occluded by vegetation/shadow at dusk. Detected via cascade architecture — OIV7 + COCO proxy catches partial features through vegetation gaps.)",

    "Tiger in deep shadow under tree canopy or at twilight":
        "Reference Image: 002711.jpg — Test Data 13\n(Tiger in deep shadow/dusk conditions under canopy. Dark enhancement (gamma correction 0.5 + CLAHE 8x8 grid) applied before detection to reveal stripe patterns.)",
}

updated = 0
for i, p in enumerate(doc.paragraphs):
    if "[Image Placeholder:" in p.text:
        for key, replacement in remaining_map.items():
            if key in p.text:
                for run in p.runs:
                    run.text = ""
                p.runs[0].text = replacement
                p.runs[0].bold = True
                p.runs[0].font.color.rgb = RGBColor(0, 51, 153)
                p.runs[0].font.size = Pt(10)
                updated += 1
                print(f"[{i}] Updated: {key[:60]}...")
                break

print(f"\nFixed {updated} remaining placeholders")

# Verify no placeholders left
remaining = 0
for p in doc.paragraphs:
    if "[Image Placeholder:" in p.text:
        remaining += 1
        print(f"  STILL REMAINING: {p.text[:100]}")

print(f"Placeholders still remaining: {remaining}")

doc.save(INPUT)
print(f"Saved to: {INPUT}")
